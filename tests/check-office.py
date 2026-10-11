"""Print a JSON summary of a .docx/.xlsx/.pptx made by the PDF → Office tools.
Used by phase2d-test.cjs. Needs python-docx, openpyxl, python-pptx."""
import json
import sys


def docx_summary(path):
    import docx
    d = docx.Document(path)
    paras = [{'text': p.text, 'style': p.style.name, 'bold': any(r.bold for r in p.runs)} for p in d.paragraphs if p.text.strip()]
    tables = [[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
    pictures = len(d.inline_shapes)
    return {'paras': paras, 'tables': tables, 'pictures': pictures}


def xlsx_summary(path):
    import openpyxl
    wb = openpyxl.load_workbook(path)
    sheets = []
    for ws in wb.worksheets:
        rows = [[c.value for c in row] for row in ws.iter_rows()]
        bold = [c.coordinate for row in ws.iter_rows() for c in row if c.value is not None and c.font and c.font.b]
        sheets.append({'name': ws.title, 'rows': rows, 'bold': bold})
    return {'sheets': sheets}


def pptx_summary(path):
    from pptx import Presentation
    p = Presentation(path)
    slides = []
    for s in p.slides:
        slides.append({
            'pictures': sum(1 for sh in s.shapes if sh.shape_type == 13),
            'text': [sh.text_frame.text for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip()],
        })
    return {'width': p.slide_width, 'height': p.slide_height, 'slides': slides}


path = sys.argv[1]
kind = path.rsplit('.', 1)[-1]
print(json.dumps({'docx': docx_summary, 'xlsx': xlsx_summary, 'pptx': pptx_summary}[kind](path), default=str))
