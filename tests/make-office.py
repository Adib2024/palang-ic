"""Make a test .docx and .xlsx (python-docx, openpyxl) for office2pdf-test.cjs.
Usage: python3 -I tests/make-office.py OUT_DIR"""
import io
import os
import struct
import sys
import zlib

import docx
import openpyxl
from docx.shared import Pt, Inches, RGBColor
from openpyxl.styles import Font

out = sys.argv[1]
os.makedirs(out, exist_ok=True)


def png(w, h, rgb):
    raw = b''.join(b'\x00' + bytes(rgb) * w for _ in range(h))
    def chunk(t, d):
        c = t + d
        return struct.pack('>I', len(d)) + c + struct.pack('>I', zlib.crc32(c) & 0xffffffff)
    return b'\x89PNG\r\n\x1a\n' + chunk(b'IHDR', struct.pack('>IIBBBBB', w, h, 8, 2, 0, 0, 0)) + chunk(b'IDAT', zlib.compress(raw)) + chunk(b'IEND', b'')


d = docx.Document()
d.add_heading('Surat Tawaran Kerja', level=1)
p = d.add_paragraph('Dengan sukacitanya kami menawarkan jawatan ')
r = p.add_run('Pegawai Kewangan')
r.bold = True
p.add_run(' kepada anda. Café — “petikan”.')
q = d.add_paragraph()
it = q.add_run('Teks condong')
it.italic = True
red = q.add_run(' merah')
red.font.color.rgb = RGBColor(0xC0, 0x10, 0x10)
d.add_heading('Faedah', level=2)
for item in ['Cuti tahunan 20 hari', 'Insurans perubatan']:
    d.add_paragraph(item, style='List Bullet')
for item in ['Lapor diri', 'Isi borang']:
    d.add_paragraph(item, style='List Number')
t = d.add_table(rows=3, cols=3)
t.style = 'Table Grid'
for i, row in enumerate([['Perkara', 'Jumlah', 'Nota'], ['Gaji', 'RM 4,500', 'Bulanan'], ['Elaun', 'RM 300', '']]):
    for j, v in enumerate(row):
        t.cell(i, j).text = v
d.add_paragraph('Gambar di bawah:')
d.add_picture(io.BytesIO(png(120, 60, (20, 120, 220))), width=Inches(2))
d.add_page_break()
long = d.add_paragraph('Halaman dua. ' + 'Ayat panjang untuk menguji balutan baris. ' * 40)
long.paragraph_format.alignment = 3  # justify
d.add_paragraph('Teks besar').runs[0].font.size = Pt(20)
d.save(os.path.join(out, 'tawaran.docx'))

wb = openpyxl.Workbook()
ws = wb.active
ws.title = 'Belanja'
ws.append(['Item', 'Jumlah', 'Tarikh', 'Kod'])
for c in ws[1]:
    c.font = Font(bold=True)
ws.append(['Sewa', 1200.5, None, '007'])
ws.append(['Elektrik', 350, None, '012'])
ws['C2'] = __import__('datetime').date(2024, 1, 31)
ws['B2'].number_format = '#,##0.00'
ws.column_dimensions['A'].width = 20
ws2 = wb.create_sheet('Ringkasan')
ws2['A1'] = 'Jumlah besar'
ws2['B1'] = 1550.5
wide = wb.create_sheet('Lebar')
wide.append([f'Lajur {i}' for i in range(1, 31)])
wb.save(os.path.join(out, 'belanja.xlsx'))
print('ok')
